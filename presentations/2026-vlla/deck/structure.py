import re, subprocess, sys, zipfile, shutil, os
SK='/root/.claude/skills/synced/79e85875-299b-4d81-bdd3-cbc6df2522d3_e0b34e4f-52d3-4c5e-8e8d-18135e62e58b/pptx/scripts'
BASES=[1,16,6,7,17,2,12,5,3,17,19,18,4,4,4,4,8,10,11,13,15,7,2]
shutil.rmtree('un',ignore_errors=True); zipfile.ZipFile('geometric.pptx').extractall('un')
new=[]
for b in BASES:
    out=subprocess.run(['python3',SK+'/add_slide.py','un',f'slide{b}.xml'],capture_output=True,text=True)
    m=re.search(r'(slide\d+\.xml)',out.stdout.split('Created')[-1] if 'Created' in out.stdout else out.stdout)
    if not m: print(out.stdout,out.stderr); sys.exit(1)
    new.append(m.group(1))
print(new)
pres=open('un/ppt/presentation.xml').read(); rels=open('un/ppt/_rels/presentation.xml.rels').read()
r2s={}
for rel in re.findall(r'<Relationship [^>]*/>',rels):
    i=re.search(r'Id="(rId\d+)"',rel).group(1); t=re.search(r'Target="([^"]+)"',rel).group(1)
    r2s[i]=t.split('/')[-1]
entries=re.findall(r'<p:sldId [^>]*/>',pres)
by={r2s[re.search(r'r:id="(rId\d+)"',e).group(1)]:e for e in entries}
lst=''.join(by[n] for n in new)
pres=re.sub(r'<p:sldIdLst>.*?</p:sldIdLst>','<p:sldIdLst>'+lst+'</p:sldIdLst>',pres,flags=re.S)
open('un/ppt/presentation.xml','w').write(pres)
print(subprocess.run(['python3',SK+'/clean.py','un'],capture_output=True,text=True).stdout[-300:])
if os.path.exists('base.pptx'): os.remove('base.pptx')
subprocess.run('cd un && zip -qXr ../base.pptx .',shell=True,check=True)
from pptx import Presentation
print(len(Presentation('base.pptx').slides))
